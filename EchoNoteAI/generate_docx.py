from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc():
    doc = Document()

    # Title
    title = doc.add_heading('EchoNote AI: Project Documentation', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Overview
    doc.add_heading('1. Overview', level=1)
    doc.add_paragraph(
        "EchoNote AI is a real-time voice-to-structured-notes assistant. It leverages Google's Gemini 1.5 Flash "
        "for natural language processing and the SpeechRecognition library for audio-to-text conversion."
    )

    # 2. System Architecture
    doc.add_heading('2. System Architecture', level=1)
    p = doc.add_paragraph()
    p.add_run('processor.py (The Engine):').bold = True
    p.add_run(" Handles the 'heavy lifting'—microphone access, speech-to-text, AI communication, and file I/O.")
    
    p2 = doc.add_paragraph()
    p2.add_run('app.py (The Orchestrator):').bold = True
    p2.add_run(" A Streamlit web interface that manages the application state (UI, buttons, and real-time updates).")

    # 3. Logic Flow Diagram
    doc.add_heading('3. Project Design Flow', level=1)
    
    flow_steps = [
        "START: Launch EchoNote AI Application",
        "  ↓",
        "USER ACTION: Click 'Start Listening' (Session state: listening = True)",
        "  ↓",
        "ENGINE: processor.listen() captures audio segments",
        "  ↓",
        "PROCESSING: Google Speech-to-Text converts audio to raw transcript",
        "  ↓",
        "USER ACTION: Click 'Stop Listening' (Session state: listening = False)",
        "  ↓",
        "AI BRAIN: processor.process_notes() sends full transcript to Gemini AI",
        "  ↓",
        "OUTPUT: AI formats notes with headers, bullet points, and TL;DR",
        "  ↓",
        "SAVE: Note is saved as a .md file and displayed in the UI sidebar"
    ]
    
    for step in flow_steps:
        p = doc.add_paragraph(step)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.name = 'Courier New'
        run.font.size = Pt(11)

    # 4. Technical Logic
    doc.add_heading('4. Technical Logic', level=1)
    
    doc.add_heading('A. The Listening Loop', level=2)
    doc.add_paragraph(
        "Uses the SpeechRecognition library. It calibrates the microphone for ambient noise for 0.5s "
        "and captures audio in small chunks (2-4 seconds) to ensure a smooth, continuous user experience."
    )

    doc.add_heading('B. Gemini AI Integration', level=2)
    doc.add_paragraph(
        "Utilizes 'gemini-1.5-flash' for speed. A custom system prompt instructs the model to act as a "
        "professional note-taker, ensuring clear formatting and a high-quality summary."
    )

    # 5. Key Libraries
    doc.add_heading('5. Libraries Used', level=1)
    libs = [
        "Streamlit: Web User Interface",
        "SpeechRecognition: Microphone interaction",
        "PyAudio: Low-level audio engine",
        "Google Generative AI SDK: Gemini integration",
        "Python-Docx: Word document generation"
    ]
    for lib in libs:
        doc.add_paragraph(lib, style='List Bullet')

    # Save the document
    doc.save('Project_Documentation.docx')
    print("Project_Documentation.docx has been created successfully!")

if __name__ == "__main__":
    create_doc()
